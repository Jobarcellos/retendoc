#!/usr/bin/env python3
"""
09_conferencia_inep.py
Conferência dos Parquet do RegDoc com as planilhas originais do Inep (IRD por escola).

O que confere (por código de escola e ano):
  - presença da escola no ano (faltantes e excedentes);
  - IRD, município, dependência administrativa, localização e nome da escola;
  - IRD municipal (municipal_consolidado.parquet) = média simples do IRD das escolas do município (todas as redes).
O que NÃO confere: ICG, AFD, IED, ATU e etapas (IN_INF, IN_FUND, IN_MED), que vêm de outros arquivos do Inep.

Uso:
  python 09_conferencia_inep.py --inep PASTA_COM_AS_PLANILHAS --repo PASTA_DO_REPOSITORIO [--anos 2025 2024] [--saida conferencia_inep.csv]

PASTA_COM_AS_PLANILHAS: pasta (com subpastas) onde estão os arquivos IRD_ESCOLAS_<ano>.xlsx já extraídos.
PASTA_DO_REPOSITORIO: pasta com escola_consolidado.parquet e municipal_consolidado.parquet.
Dependências: pandas, pyarrow, openpyxl.
"""
import argparse, glob, os, re, sys
import pandas as pd
import openpyxl

COLS = ["ANO", "REGIAO", "UF", "CO_MUNICIPIO", "NO_MUNICIPIO", "CO_ENTIDADE",
        "NO_ENTIDADE", "LOCALIZACAO", "DEPENDENCIA", "IRD"]
DEP = {1: "Federal", 2: "Estadual", 3: "Municipal", 4: "Privada"}


def ler_planilha_escolas(caminho):
    """Lê a planilha do Inep por posição de coluna (o cabeçalho muda entre anos)."""
    ws = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ws = ws[ws.sheetnames[0]]
    linhas = [r[:10] for r in ws.iter_rows(values_only=True)
              if r and isinstance(r[0], (int, float)) and isinstance(r[5], (int, float))]
    d = pd.DataFrame(linhas, columns=COLS)
    for c in ["REGIAO", "UF", "NO_MUNICIPIO", "NO_ENTIDADE", "LOCALIZACAO", "DEPENDENCIA"]:
        d[c] = d[c].astype(str).str.strip()
    d["IRD"] = pd.to_numeric(d["IRD"], errors="coerce")
    for c in ["ANO", "CO_MUNICIPIO", "CO_ENTIDADE"]:
        d[c] = d[c].astype("int64")
    return d


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--inep", required=True)
    ap.add_argument("--repo", required=True)
    ap.add_argument("--anos", nargs="*", type=int)
    ap.add_argument("--saida", default="conferencia_inep.csv")
    a = ap.parse_args()

    esc = pd.read_parquet(os.path.join(a.repo, "escola_consolidado.parquet"))
    mun = pd.read_parquet(os.path.join(a.repo, "municipal_consolidado.parquet"))
    for d, cols in [(esc, ["CO_ENTIDADE", "CO_MUNICIPIO", "ANO"]), (mun, ["CO_MUNICIPIO", "ANO"])]:
        for c in cols:
            d[c] = d[c].astype("int64")

    arquivos = {}
    for p in glob.glob(os.path.join(a.inep, "**", "*ESCOLAS_*.xlsx"), recursive=True):
        m = re.search(r"ESCOLAS_(20\d\d)\.xlsx$", p)
        if m:
            arquivos[int(m.group(1))] = p
    anos = sorted(a.anos) if a.anos else sorted(arquivos)
    faltam = [x for x in anos if x not in arquivos]
    if faltam:
        sys.exit(f"Planilhas não encontradas para: {faltam}")

    saida = []
    for ano in anos:
        i = ler_planilha_escolas(arquivos[ano])
        q = esc[esc.ANO == ano]
        m = q.merge(i, on="CO_ENTIDADE", how="outer", suffixes=("_r", "_i"), indicator=True)
        b = m[m._merge == "both"]
        dep_r = b.TP_DEPENDENCIA.map(DEP)
        loc_r = b.TP_LOCALIZACAO.astype(str).str.strip().str.lower()
        g = i.groupby("CO_MUNICIPIO").IRD.mean()
        mm = mun[mun.ANO == ano].set_index("CO_MUNICIPIO").IRD
        j = pd.concat([mm, g], axis=1, keys=["p", "c"])
        saida.append(dict(
            ano=ano, registros_inep=len(i), registros_regdoc=len(q),
            so_inep=int((m._merge == "right_only").sum()), so_regdoc=int((m._merge == "left_only").sum()),
            ird_difere=int(((b.IRD_r - b.IRD_i).abs() > 1e-6).sum()),
            ird_dif_max=float((b.IRD_r - b.IRD_i).abs().max()),
            municipio_difere=int((b.CO_MUNICIPIO_r != b.CO_MUNICIPIO_i).sum()),
            dependencia_difere=int((dep_r != b.DEPENDENCIA).sum()),
            localizacao_difere=int((loc_r != b.LOCALIZACAO.str.lower()).sum()),
            nome_difere=int((b.NO_ENTIDADE_r.astype(str).str.strip().str.upper()
                             != b.NO_ENTIDADE_i.str.upper()).sum()),
            municipios_regdoc=int(j.p.notna().sum()), municipios_inep=int(j.c.notna().sum()),
            ird_municipal_difere=int(((j.p - j.c).abs() > 1e-6).sum() + (j.p.isna() ^ j.c.isna()).sum()),
            media_escolas=round(float(i.IRD.mean()), 4), media_municipios=round(float(mm.mean()), 4)))
        print(saida[-1], flush=True)
    pd.DataFrame(saida).to_csv(a.saida, index=False)
    print("Gravado:", a.saida)


if __name__ == "__main__":
    main()
